"""Genera el 'Reporte de horas de servicio cívico' a partir de la plantilla de septiembre."""
import io
import math
from copy import copy

from openpyxl import load_workbook
from openpyxl.worksheet.pagebreak import Break

from kobo import MESES

PRIMERA_FILA = 12      # desde aquí la plantilla se reconstruye según el calendario del mes
COLS = "ABCDEFGHIJK"

# Filas de la plantilla (Mes_de_septiembre) que sirven de modelo para cada parte
P_ENCAB = (12, 13)     # encabezado de semana
P_DIA = 29             # fila de día (como el ejemplo ya llenado de la plantilla)
P_SUBTOTAL = 18
P_BLANCO = 51
P_TOTAL = 53
P_FIRMA = (63, 64)
ALTO_DIA = 35.25


def _modelo(ws, fila, con_valores=True):
    return {"alto": ws.row_dimensions[fila].height,
            "celdas": {c: (copy(ws[f"{c}{fila}"]._style), ws[f"{c}{fila}"].value if con_valores else None)
                       for c in COLS}}


def _escribir_fila(ws, fila, modelo, valores=None, alto=None, estilos_extra=None):
    for c in COLS:
        estilo, valor = modelo["celdas"][c]
        if estilos_extra and c in estilos_extra:
            estilo = estilos_extra[c]
        celda = ws[f"{c}{fila}"]
        celda._style = copy(estilo)
        celda.value = valor
    for c, v in (valores or {}).items():
        ws[f"{c}{fila}"].value = v
    ws.row_dimensions[fila].height = alto if alto is not None else modelo["alto"]


def _alto_para(texto_c, texto_f):
    """Alto de fila para que quepa el texto (C:E combinadas ~50 caracteres, F ~11)."""
    def lineas(texto, ancho):
        return sum(max(1, math.ceil(len(l) / ancho)) for l in str(texto).split("\n")) if texto else 1
    n = max(lineas(texto_c, 50), lineas(texto_f, 11))
    return max(ALTO_DIA, 13.0 * n + 8)


def generar_reporte(plantilla: str, datos: dict) -> bytes:
    wb = load_workbook(plantilla)
    ws = wb.worksheets[0]

    m = {
        "encab": [_modelo(ws, f) for f in P_ENCAB],
        "dia": _modelo(ws, P_DIA, con_valores=False),
        "subtotal": _modelo(ws, P_SUBTOTAL),
        "blanco": _modelo(ws, P_BLANCO, con_valores=False),
        "total": _modelo(ws, P_TOTAL),
        "firma": [_modelo(ws, f) for f in P_FIRMA],
    }
    # Bordes de las celdas combinadas Semana (A) y Firma (K): arriba / en medio / abajo
    borde = {pos: {c: copy(ws[f"{c}{f}"]._style) for c in "AK"}
             for pos, f in (("medio", 15), ("abajo", 17))}

    # --- borrar todo desde la fila 12 hacia abajo
    for rango in list(ws.merged_cells.ranges):
        if rango.max_row >= PRIMERA_FILA:
            ws.unmerge_cells(str(rango))
    for (f, c) in [k for k in ws._cells if k[0] >= PRIMERA_FILA]:
        del ws._cells[(f, c)]
    for f in [k for k in ws.row_dimensions if k >= PRIMERA_FILA]:
        del ws.row_dimensions[f]

    # --- encabezado
    ws["C2"] = datos["nombre"]
    ws["C5"] = datos["departamento"]
    ws["C8"] = datos["municipio"]
    ws["C10"] = f"{MESES[datos['mes'] - 1]} {datos['anio']}"
    ws.title = MESES[datos["mes"] - 1].upper()

    fila, subtotales = PRIMERA_FILA, []
    for i, semana in enumerate(datos["semanas"]):
        if i == 4:  # como en la plantilla: dos filas en blanco antes de la 5.ª semana (salto de página)
            for _ in range(2):
                _escribir_fila(ws, fila, m["blanco"], alto=16.5)
                fila += 1
            ws.row_breaks.append(Break(id=fila - 1))   # la 5.ª semana empieza en la página 2
        # encabezado de la semana (2 filas)
        h = fila
        _escribir_fila(ws, h, m["encab"][0])
        _escribir_fila(ws, h + 1, m["encab"][1])
        for rango in (f"A{h}:A{h+1}", f"B{h}:B{h+1}", f"C{h}:E{h+1}", f"F{h}:F{h+1}",
                      f"G{h}:G{h+1}", f"H{h}:I{h}", f"J{h}:J{h+1}", f"K{h}:K{h+1}"):
            ws.merge_cells(rango)
        fila += 2

        primera = fila
        n = len(semana["dias"])
        for k, dia in enumerate(semana["dias"]):
            modelo = m["dia"]
            extra = None if k == 0 else (borde["abajo"] if k == n - 1 else borde["medio"])
            valores = {"B": dia["fecha"]}
            if k == 0:
                valores["A"] = semana["numero"]
            alto = ALTO_DIA
            if dia["estado"] == "trabajo":
                valores.update({"C": dia["actividades"] or None, "F": dia["comunidades"] or None,
                                "G": dia["familias"], "H": 8, "I": 16, "J": 8})
                alto = _alto_para(dia["actividades"], dia["comunidades"])
            elif dia["estado"] in ("feriado", "no_laborado"):
                valores.update({"C": dia["actividades"], "F": 0, "G": 0, "H": 0, "I": 0, "J": 0})
            _escribir_fila(ws, fila, modelo, valores, alto=alto, estilos_extra=extra)
            ws.merge_cells(f"C{fila}:E{fila}")
            fila += 1
        ultima = fila - 1
        if n > 1:
            ws.merge_cells(f"A{primera}:A{ultima}")
            ws.merge_cells(f"K{primera}:K{ultima}")

        _escribir_fila(ws, fila, m["subtotal"], {"J": f"=SUM(J{primera}:J{ultima})"})
        ws.merge_cells(f"F{fila}:I{fila}")
        subtotales.append(f"J{fila}")
        fila += 1

    # --- total y firma (mismas distancias que en la plantilla)
    for _ in range(2):
        _escribir_fila(ws, fila, m["blanco"]); fila += 1
    _escribir_fila(ws, fila, m["total"], {"D": "=" + "+".join(subtotales)})
    fila += 1
    for _ in range(P_FIRMA[0] - P_TOTAL - 1):
        _escribir_fila(ws, fila, m["blanco"]); fila += 1
    for k in range(2):
        _escribir_fila(ws, fila, m["firma"][k])
        ws.merge_cells(f"A{fila}:K{fila}")
        fila += 1

    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
